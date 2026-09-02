#!/usr/bin/env python3
"""Generate the journal-format Word manuscript and its PDF rendering."""
from __future__ import annotations

import subprocess
from pathlib import Path

from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[2]
SRC = ROOT / "manuscript/drafts/STANDARD_MANUSCRIPT.md"
OUT_DOCX = ROOT / "outputs/submission/STANDARD_MANUSCRIPT_JH_FORMATTED.docx"
OUT_PDF = ROOT / "outputs/submission/STANDARD_MANUSCRIPT_JH_FORMATTED.pdf"

FONT_NAME = "Times New Roman"
FONT_SIZE = Pt(12)


def enable_line_numbers(doc: Document) -> None:
    """Add reviewer-friendly line numbers that run continuously through the manuscript."""
    for section in doc.sections:
        sect_pr = section._sectPr
        existing = sect_pr.find(qn("w:lnNumType"))
        if existing is not None:
            sect_pr.remove(existing)
        line_nums = OxmlElement("w:lnNumType")
        line_nums.set(qn("w:countBy"), "1")
        line_nums.set(qn("w:distance"), "360")
        line_nums.set(qn("w:restart"), "continuous")
        sect_pr.append(line_nums)


def add_page_numbers(doc: Document) -> None:
    """Add a centered, editable PAGE field to each section footer."""
    for section in doc.sections:
        footer = section.footer
        paragraph = footer.paragraphs[0]
        paragraph.alignment = 1
        paragraph.clear()
        p_pr = paragraph._p.get_or_add_pPr()
        suppress = OxmlElement("w:suppressLineNumbers")
        p_pr.append(suppress)
        run = paragraph.add_run()
        run.font.name = FONT_NAME
        run.font.size = Pt(9)
        begin = OxmlElement("w:fldChar")
        begin.set(qn("w:fldCharType"), "begin")
        instr = OxmlElement("w:instrText")
        instr.set(qn("xml:space"), "preserve")
        instr.text = " PAGE "
        separate = OxmlElement("w:fldChar")
        separate.set(qn("w:fldCharType"), "separate")
        text = OxmlElement("w:t")
        text.text = "1"
        end = OxmlElement("w:fldChar")
        end.set(qn("w:fldCharType"), "end")
        run._r.extend([begin, instr, separate, text, end])


def set_three_line_tables(doc: Document) -> None:
    """Use journal-style top, header-separator, and bottom rules without verticals."""
    for table in doc.tables:
        table.autofit = True
        tbl_pr = table._tbl.tblPr
        existing = tbl_pr.find(qn("w:tblBorders"))
        if existing is not None:
            tbl_pr.remove(existing)
        borders = OxmlElement("w:tblBorders")
        for edge, size in (("top", "12"), ("bottom", "12")):
            tag = OxmlElement(f"w:{edge}")
            tag.set(qn("w:val"), "single")
            tag.set(qn("w:sz"), size)
            tag.set(qn("w:space"), "0")
            tag.set(qn("w:color"), "000000")
            borders.append(tag)
        for edge in ("left", "insideV", "right"):
            tag = OxmlElement(f"w:{edge}")
            tag.set(qn("w:val"), "nil")
            borders.append(tag)
        tbl_pr.append(borders)
        if table.rows:
            for cell in table.rows[0].cells:
                tc_pr = cell._tc.get_or_add_tcPr()
                header_border = OxmlElement("w:tcBorders")
                header_bottom = OxmlElement("w:bottom")
                header_bottom.set(qn("w:val"), "single")
                header_bottom.set(qn("w:sz"), "8")
                header_bottom.set(qn("w:space"), "0")
                header_bottom.set(qn("w:color"), "000000")
                header_border.append(header_bottom)
                tc_pr.append(header_border)
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.bold = True


def main() -> None:
    # Generate the base DOCX from the canonical Markdown source.
    subprocess.run([
        "pandoc", str(SRC), "-o", str(OUT_DOCX),
        "--standalone",
        "--resource-path=manuscript/drafts",
    ], check=True, cwd=ROOT)

    # Apply journal-style formatting.
    doc = Document(str(OUT_DOCX))
    normal = doc.styles["Normal"]
    normal.font.name = FONT_NAME
    normal.font.size = FONT_SIZE
    pf = normal.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.space_after = Pt(6)

    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    enable_line_numbers(doc)
    add_page_numbers(doc)
    set_three_line_tables(doc)

    doc.save(str(OUT_DOCX))

    # Convert the formatted document to PDF.
    subprocess.run([
        "soffice", "--headless", "--convert-to", "pdf",
        "--outdir", str(OUT_DOCX.parent), str(OUT_DOCX),
    ], check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)

    print("saved", OUT_DOCX, OUT_PDF)


if __name__ == "__main__":
    main()
